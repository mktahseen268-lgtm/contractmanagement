"""Repository and report exports.

The brief asks for XLSX rather than CSV, and the reason is not cosmetic: these go to auditors.
A CSV loses every date type and turns reference numbers into scientific notation the moment
somebody opens it in Excel.

The tests read the workbook back with `openpyxl` rather than asserting on bytes — the claim is
that the file *is a spreadsheet with the right values in the right cells*, and a byte assertion
would pass for a file Excel refuses to open.

Requirements: SOW-31.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid

import pytest
from openpyxl import load_workbook

from app import export_service, models, repository_service
from app.export_service import ExportError


@pytest.fixture()
def workspace(db, make_user):
    owner, tenant = make_user(email=f"exp-{uuid.uuid4().hex[:8]}@example.com", name="Owner")
    return {"tenant": tenant, "owner": owner}


def _contract(db, ws, **kwargs) -> models.Contract:
    c = models.Contract(
        tenant_id=ws["tenant"].id,
        reference_no=kwargs.pop("reference_no", f"CM-{uuid.uuid4().hex[:6]}"),
        title=kwargs.pop("title", "Vendor Services Agreement"),
        type=kwargs.pop("type", "vendor"), status=kwargs.pop("status", "active"),
        owner_id=ws["owner"].id, created_by=ws["owner"].id,
        counterparty=kwargs.pop("counterparty", "Acme Trading"),
        value=kwargs.pop("value", 1_000_000), currency="PKR",
        effective_date=kwargs.pop("effective_date", dt.date(2026, 1, 1)),
        end_date=kwargs.pop("end_date", dt.date(2026, 12, 31)),
        **kwargs,
    )
    db.add(c)
    db.flush()
    return c


def _open(payload: bytes):
    return load_workbook(io.BytesIO(payload))


# ---------------------------------------------------------------------------------------


def test_the_workbook_is_a_real_spreadsheet(db, workspace):
    _contract(db, workspace)
    payload, count = export_service.build_workbook(db, workspace["tenant"].id,
                                                   "contracts", {})
    book = _open(payload)
    assert count == 1
    assert book.sheetnames == ["Filters", "Contracts"]


def test_the_first_sheet_records_what_the_export_actually_is(db, workspace):
    """An export whose provenance is not on its own face is one nobody can defend later."""
    _contract(db, workspace, status="active")
    payload, _ = export_service.build_workbook(
        db, workspace["tenant"].id, "contracts", {"status": ["active"], "q": "vendor"},
        generated_by="Aisha Khan", org_name="Mobilink Microfinance Bank")

    cover = _open(payload)["Filters"]
    text = "\n".join(
        str(cell.value) for row in cover.iter_rows() for cell in row if cell.value is not None
    )
    assert "Mobilink Microfinance Bank" in text
    assert "Aisha Khan" in text
    assert "active" in text
    assert "vendor" in text


def test_an_unfiltered_export_says_so_rather_than_leaving_the_sheet_blank(db, workspace):
    _contract(db, workspace)
    payload, _ = export_service.build_workbook(db, workspace["tenant"].id, "contracts", {})
    text = "\n".join(
        str(c.value) for row in _open(payload)["Filters"].iter_rows() for c in row
        if c.value is not None
    )
    assert "whole repository" in text


def test_dates_stay_dates_and_numbers_stay_numbers(db, workspace):
    """The reason this is XLSX and not CSV."""
    _contract(db, workspace, value=1_250_000, effective_date=dt.date(2026, 3, 1))
    payload, _ = export_service.build_workbook(db, workspace["tenant"].id, "contracts", {})
    sheet = _open(payload)["Contracts"]
    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    row = [c.value for c in next(sheet.iter_rows(min_row=2, max_row=2))]

    assert isinstance(row[headers.index("Value")], (int, float))
    assert isinstance(row[headers.index("Effective")], (dt.date, dt.datetime))


def test_a_reference_number_survives_as_text(db, workspace):
    """The classic CSV failure: a reference opened in Excel as a number."""
    _contract(db, workspace, reference_no="C-2026-0007")
    payload, _ = export_service.build_workbook(db, workspace["tenant"].id, "contracts", {})
    sheet = _open(payload)["Contracts"]
    assert next(sheet.iter_rows(min_row=2, max_row=2))[0].value == "C-2026-0007"


def test_filters_narrow_the_export(db, workspace):
    _contract(db, workspace, status="active")
    _contract(db, workspace, status="draft")
    _payload, count = export_service.build_workbook(
        db, workspace["tenant"].id, "contracts", {"status": ["active"]})
    assert count == 1


def test_obligations_can_be_exported(db, workspace):
    contract = _contract(db, workspace)
    db.add(models.Obligation(
        tenant_id=workspace["tenant"].id, contract_id=contract.id,
        title="Quarterly service report", due_date=dt.date(2026, 6, 30),
        created_by=workspace["owner"].id))
    db.flush()

    payload, count = export_service.build_workbook(db, workspace["tenant"].id,
                                                   "obligations", {})
    assert count == 1
    sheet = _open(payload)["Obligations"]
    assert "Quarterly service report" in [c.value for c in next(sheet.iter_rows(min_row=2, max_row=2))]


def test_parties_export_counts_their_agreements(db, workspace):
    party = repository_service.create_party(db, workspace["tenant"].id,
                                            {"name": "Acme Trading", "registration_no": "123"},
                                            actor=workspace["owner"])
    contract = _contract(db, workspace)
    repository_service.link_contract_to_party(db, contract, party)
    db.flush()

    payload, count = export_service.build_workbook(db, workspace["tenant"].id, "parties", {})
    assert count == 1
    sheet = _open(payload)["Parties"]
    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    row = [c.value for c in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert row[headers.index("Agreements")] == 1


def test_the_audit_log_can_be_exported_in_chain_order(db, workspace):
    """An auditor reads it in the order the chain was built, not by timestamp."""
    from app import audit

    for action in ("a", "b", "c"):
        audit.record(db, tenant_id=workspace["tenant"].id, action=action, actor=None,
                     object_type="test", object_label=action)
    db.flush()

    payload, count = export_service.build_workbook(db, workspace["tenant"].id, "audit", {})
    sheet = _open(payload)["Audit"]
    sequences = [row[1].value for row in sheet.iter_rows(min_row=2)]
    assert count >= 3
    assert sequences == sorted(sequences)


def test_another_tenants_rows_are_never_exported(db, workspace, make_user):
    other_owner, other_tenant = make_user(email=f"o-{uuid.uuid4().hex[:8]}@example.com",
                                          name="Other")
    db.add(models.Contract(
        tenant_id=other_tenant.id, reference_no="X-1", title="Theirs", type="vendor",
        status="active", owner_id=other_owner.id, created_by=other_owner.id, currency="PKR"))
    db.flush()
    _payload, count = export_service.build_workbook(db, workspace["tenant"].id,
                                                    "contracts", {})
    assert count == 0


def test_an_unknown_export_kind_is_refused(db, workspace):
    with pytest.raises(ExportError, match="Nothing called"):
        export_service.build_workbook(db, workspace["tenant"].id, "everything", {})


def test_an_empty_repository_still_produces_a_usable_workbook(db, workspace):
    """A zero-row export must open, not fail — "there is nothing" is a valid answer."""
    payload, count = export_service.build_workbook(db, workspace["tenant"].id,
                                                   "contracts", {})
    assert count == 0
    sheet = _open(payload)["Contracts"]
    assert [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))][0] == "Reference"


def test_running_an_export_records_a_downloadable_job(db, workspace):
    _contract(db, workspace)
    job = export_service.run_export(db, workspace["tenant"].id, "contracts", {},
                                    actor=workspace["owner"], org_name="MMBL")

    assert job.status == "succeeded"
    assert job.progress == 100
    assert job.href == f"/exports/{job.id}/download"
    assert "1 row" in job.result_summary


def test_the_generated_file_is_actually_in_storage(db, workspace):
    from app.storage import get_storage

    _contract(db, workspace)
    job = export_service.run_export(db, workspace["tenant"].id, "contracts", {},
                                    actor=workspace["owner"])
    key = export_service.storage_key(job)
    assert get_storage().exists(key)
    with get_storage().open_stream(key) as stream:
        assert _open(stream.read()).sheetnames == ["Filters", "Contracts"]


def test_a_failed_export_leaves_a_visible_failed_job(db, workspace):
    """A failure the user cannot see is worse than a failure."""
    with pytest.raises(ExportError):
        export_service.run_export(db, workspace["tenant"].id, "nonsense", {},
                                  actor=workspace["owner"])
    job = db.query(models.BackgroundJob).filter_by(
        tenant_id=workspace["tenant"].id, status="failed").one()
    assert job.error


def test_the_filename_names_the_export_and_the_day(db, workspace):
    _contract(db, workspace)
    job = export_service.run_export(db, workspace["tenant"].id, "contracts", {},
                                    actor=workspace["owner"])
    name = export_service.filename(job)
    assert name.startswith("contracts-")
    assert name.endswith(".xlsx")
